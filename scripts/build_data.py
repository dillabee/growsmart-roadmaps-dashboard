"""
Build dashboard JSON from the Dashboard Ready workbook.

Usage:
    python scripts/build_data.py [path-to-xlsx]

Reads "Dashboard Data - Community Profiles - NSP GSM.xlsx" and writes
public/data/dashboard_data.json. Run this whenever the workbook changes
(annual updates in years 2 and 3), then rebuild the site.

Validation: fails loudly on missing sheets/columns; warns on missing
community/year combinations.
"""
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

DEFAULT_XLSX = (
    r"C:\Users\Dillon Roberts\OneDrive - Camoin Associates\Clients"
    r"\North Star Planning\GrowSmart Maine\04 Dashboard Ready Files"
    r"\Dashboard Data - Community Profiles - NSP GSM.xlsx"
)

COMMUNITIES = ["Calais", "Danforth", "Houlton", "Machias", "Presque Isle", "Van Buren"]
COMPARISONS = ["Aroostook County", "Washington County", "Maine"]
REGION = "Roadmaps Region"

# Ratio indicators: id -> (numerator vars, denominator vars)
RATES = {
    "unemploymentRate": (["Unemployed"], ["Labor Force"]),
    "lfpr": (["Labor Force"], ["Population Age 16+"]),
    "pctBachelors": (["Bachelors or Higher"], ["Education Total"]),
    "pctLessHS": (["Less than High School"], ["Education Total"]),
    "homeownershipRate": (["Owner Occupied"], ["Tenure Total"]),
    "renterCostBurden": (["Renter Cost >30%"], ["Renter Cost Burden Total"]),
    "ownerMortgageCostBurden": (["Owner Mortgage 30%+"], ["Owner Mortgage Total"]),
    "pctWFH": (["Work From Home"], ["Total Workers"]),
}
# Level indicators copied straight through (additive ones roll up to region)
LEVELS_ADDITIVE = {
    "population": "Total Population",
    "households": "Total Households",
    "laborForce": "Labor Force",
}
LEVELS_NON_ADDITIVE = {
    "medianAge": "Median Age",
    "medianHHI": "Median Household Income",
}
RACE_VARS = {
    "White": "White", "Black": "Black", "AIAN": "AIAN", "Asian": "Asian",
    "NHPI": "NHPI", "Other": "Other", "Two+ races": "Multi Not Hispanic",
    "Hispanic": "Hispanic ",
}
INCOME_BRACKETS = [
    "<$10,000", "$10,000 - $14,999", "$15,000 - $19,999", "$20,000 - $24,999",
    "$25,000 - $29,999", "$30,000 - $34,999", "$35,000 - $39,999",
    "$40,000 - $44,999", "$45,000 - $49,999", "$50,000 - $59,999",
    "$60,000 - $74,999", "$75,000 - $99,999", "$100,000 - $124,999",
    "$125,000 - $149,999", "$150,000 - $199,999", "$200,000+",
]


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def rows_of(ws):
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    for r in it:
        if any(v is not None for v in r):
            yield dict(zip(header, r))


def build_demographics(ws):
    # cell[geo][year][var] = summed estimate
    cell = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    for r in rows_of(ws):
        geo, var, year = r["Community"], r["Clean Var Name"], str(r["year"])
        v = num(r["estimate"])
        if geo is None or var is None or v is None:
            continue
        cell[geo][year][var] += v

    years = sorted({y for g in cell.values() for y in g})
    for c in COMMUNITIES:
        for y in years:
            if not cell.get(c, {}).get(y):
                print(f"WARNING: no demographic data for {c} {y}")

    # Region rollup: sum the six communities per year/var
    for y in years:
        agg = defaultdict(float)
        for c in COMMUNITIES:
            for var, v in cell[c][y].items():
                agg[var] += v
        cell[REGION][y] = agg

    out = {}
    for geo in COMMUNITIES + [REGION] + COMPARISONS:
        g = {}
        for ind, var in LEVELS_ADDITIVE.items():
            g[ind] = {y: cell[geo][y].get(var) for y in years}
        for ind, var in LEVELS_NON_ADDITIVE.items():
            if geo == REGION:
                g[ind] = {y: None for y in years}  # medians don't sum
            else:
                g[ind] = {y: cell[geo][y].get(var) for y in years}
        for ind, (nums, dens) in RATES.items():
            g[ind] = {}
            for y in years:
                n = sum(cell[geo][y].get(v, 0) for v in nums)
                d = sum(cell[geo][y].get(v, 0) for v in dens)
                g[ind][y] = round(n / d * 100, 1) if d else None
        g["race"] = {
            y: {
                label: cell[geo][y].get(var)
                for label, var in RACE_VARS.items()
            }
            for y in years
        }
        g["incomeDist"] = {
            y: {b: cell[geo][y].get(b) for b in INCOME_BRACKETS} for y in years
        }
        out[geo] = g
    return out, years


def build_econ(ws):
    out = defaultdict(lambda: {"totals": {}, "industries": defaultdict(dict)})
    for r in rows_of(ws):
        geo = r["Area Name"]
        year = str(r["Year"])
        rec = {
            "establishments": num(r["Establishments"]),
            "employment": num(r["Average Employment"]),
            "totalWages": num(r["Total Wages"]),
            "avgWeeklyWage": num(r["Average Weekly Wage"]),
        }
        title = str(r["NAICS Title"]).strip()
        if str(r["NAICS"]) == "0" or title in ("Total, All Industries", "Total All Industries"):
            out[geo]["totals"][year] = rec
        else:
            out[geo]["industries"][title][year] = rec
    # Region rollup: sum the six communities (avg weekly wage derived from
    # summed wages / summed employment). Suppressed industries are treated
    # as zero, so regional industry values are a floor, not an exact sum.
    region = {"totals": {}, "industries": defaultdict(dict)}

    def summed(records):
        est = sum(r["establishments"] or 0 for r in records)
        emp = sum(r["employment"] or 0 for r in records)
        wages = sum(r["totalWages"] or 0 for r in records)
        return {
            "establishments": est or None,
            "employment": emp or None,
            "totalWages": wages or None,
            "avgWeeklyWage": round(wages / emp / 52) if emp else None,
        }

    years = {y for c in COMMUNITIES for y in out.get(c, {}).get("totals", {})}
    for y in sorted(years):
        recs = [out[c]["totals"][y] for c in COMMUNITIES if y in out.get(c, {}).get("totals", {})]
        if recs:
            region["totals"][y] = summed(recs)
        titles = {t for c in COMMUNITIES for t in out.get(c, {}).get("industries", {})}
        for t in titles:
            recs = [
                out[c]["industries"][t][y]
                for c in COMMUNITIES
                if y in out.get(c, {}).get("industries", {}).get(t, {})
            ]
            if recs:
                region["industries"][t][y] = summed(recs)
    out[REGION] = region

    return {
        geo: {
            "totals": d["totals"],
            "industries": [
                {"title": t, **{f"y{y}": v for y, v in yrs.items()}}
                for t, yrs in sorted(d["industries"].items())
            ],
        }
        for geo, d in out.items()
    }


def build_trade(ws):
    trend = defaultdict(dict)  # port -> year -> {exports, imports}
    detail = defaultdict(lambda: defaultdict(float))  # (port,type,dim) -> name -> value
    LATEST = "2025"
    for r in rows_of(ws):
        port = str(r["Port"]).replace(", ME (Port)", "")
        ttype, commodity, country = r["Type"], str(r["Commodity"]).strip(), str(r["Country"]).strip()
        year, v = str(r["Time"]), num(r[[k for k in r if k.startswith("Total")][0]])
        if v is None:
            continue
        if commodity == "All Commodities" and country == "World Total":
            trend[port].setdefault(year, {})[ttype.lower() + "s"] = v
        if year == LATEST:
            if country == "World Total" and commodity != "All Commodities":
                detail[(port, ttype, "commodity")][commodity] += v
            if commodity == "All Commodities" and country != "World Total":
                detail[(port, ttype, "country")][country] += v

    out = {}
    for port in trend:
        top = {}
        for ttype in ("Export", "Import"):
            for dim in ("commodity", "country"):
                items = sorted(
                    detail[(port, ttype, dim)].items(), key=lambda kv: -kv[1]
                )[:8]
                top[f"{ttype.lower()}_{dim}"] = [
                    {"name": k, "value": v} for k, v in items
                ]
        out[port] = {"trend": trend[port], "top": top, "topYear": LATEST}
    return out


def build_cbp(ws):
    # port -> "Measure / Category" -> FY -> summed count
    # The source sheet contains some exact duplicate rows (e.g. Houlton
    # FY2022), so identical rows are counted once.
    out = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    seen = set()
    dupes = 0
    for r in rows_of(ws):
        key = tuple(sorted((str(k), str(v)) for k, v in r.items()))
        if key in seen:
            dupes += 1
            continue
        seen.add(key)
        port = str(r["Port of Entry"]).split(",")[0].title()
        key = f'{r["Measure Name"]} / {r["Category"]}'
        v = num(r["Count"])
        if v is None:
            continue
        out[port][key][str(r["FY"])] += v
    if dupes:
        print(f"WARNING: skipped {dupes} exact duplicate CBP rows in source sheet")
    return {p: {k: dict(fy) for k, fy in m.items()} for p, m in out.items()}


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(DEFAULT_XLSX)
    if not xlsx.exists():
        sys.exit(f"ERROR: workbook not found: {xlsx}")
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    required = ["Demographic Socioeconomic", "Economic Base", "Trade", "Customs and Border Patrol"]
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        sys.exit(f"ERROR: missing sheets: {missing}")

    demo, demo_years = build_demographics(wb["Demographic Socioeconomic"])
    data = {
        "meta": {
            "communities": COMMUNITIES,
            "comparisons": COMPARISONS,
            "region": REGION,
            "demoYears": demo_years,
            "sourceFile": xlsx.name,
        },
        "demographics": demo,
        "econ": build_econ(wb["Economic Base"]),
        "trade": build_trade(wb["Trade"]),
        "cbp": build_cbp(wb["Customs and Border Patrol"]),
    }
    out = Path(__file__).resolve().parent.parent / "public" / "data" / "dashboard_data.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data), encoding="utf-8")
    print(f"Wrote {out} ({out.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
